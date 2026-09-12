"""
Document loading and text chunking for AI Boardroom RAG.

Supported formats:
- PDF
- DOCX
- PPTX
- TXT
"""

from pathlib import Path
from typing import List, Dict

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


class DocumentLoader:
    """Load supported business documents and split them into chunks."""

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".pptx",
        ".txt",
    }

    def __init__(
        self,
        chunk_size: int = 1200,
        chunk_overlap: int = 200,
    ):
        if chunk_size <= 0:
            raise ValueError("chunk_size must be positive")

        if chunk_overlap < 0:
            raise ValueError("chunk_overlap cannot be negative")

        if chunk_overlap >= chunk_size:
            raise ValueError(
                "chunk_overlap must be smaller than chunk_size"
            )

        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    # =========================================================
    # PUBLIC API
    # =========================================================

    def load(self, file_path: str) -> List[Dict]:
        """
        Load a document and return text chunks.

        Each chunk contains:
        - text
        - source
        - page/slide metadata when available
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Document not found: {file_path}"
            )

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported document type: {extension}. "
                f"Supported types: "
                f"{', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )

        if extension == ".pdf":
            pages = self._load_pdf(path)

        elif extension == ".docx":
            pages = self._load_docx(path)

        elif extension == ".pptx":
            pages = self._load_pptx(path)

        else:
            pages = self._load_txt(path)

        return self._create_chunks(
            pages,
            source_name=path.name,
        )

    # =========================================================
    # PDF
    # =========================================================

    def _load_pdf(self, path: Path) -> List[Dict]:
        reader = PdfReader(str(path))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            text = self._clean_text(text)

            if text:
                pages.append(
                    {
                        "text": text,
                        "page": page_number,
                    }
                )

        return pages

    # =========================================================
    # DOCX
    # =========================================================

    def _load_docx(self, path: Path) -> List[Dict]:
        document = Document(str(path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = self._clean_text(
                paragraph.text
            )

            if text:
                paragraphs.append(text)

        # Include table content as well.
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:
                    cell_text = self._clean_text(
                        cell.text
                    )

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    paragraphs.append(
                        " | ".join(cells)
                    )

        text = "\n".join(paragraphs)

        if not text:
            return []

        return [
            {
                "text": text,
                "page": None,
            }
        ]

    # =========================================================
    # PPTX
    # =========================================================

    def _load_pptx(self, path: Path) -> List[Dict]:
        presentation = Presentation(str(path))

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            texts = []

            for shape in slide.shapes:

                if not hasattr(shape, "text"):
                    continue

                text = self._clean_text(
                    shape.text
                )

                if text:
                    texts.append(text)

            slide_text = "\n".join(texts)

            if slide_text:
                slides.append(
                    {
                        "text": slide_text,
                        "page": slide_number,
                    }
                )

        return slides

    # =========================================================
    # TXT
    # =========================================================

    def _load_txt(self, path: Path) -> List[Dict]:

        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        text = self._clean_text(text)

        if not text:
            return []

        return [
            {
                "text": text,
                "page": None,
            }
        ]

    # =========================================================
    # TEXT CLEANING
    # =========================================================

    @staticmethod
    def _clean_text(text: str) -> str:

        if not text:
            return ""

        lines = []

        for line in text.splitlines():

            line = " ".join(
                line.split()
            )

            if line:
                lines.append(line)

        return "\n".join(lines).strip()

    # =========================================================
    # CHUNKING
    # =========================================================

    def _create_chunks(
        self,
        pages: List[Dict],
        source_name: str,
    ) -> List[Dict]:
        """
        Create overlapping character-based chunks.

        Character-based chunking keeps the implementation
        deterministic and lightweight.
        """

        chunks = []

        chunk_id = 0

        for page_data in pages:

            text = page_data["text"]
            page = page_data.get("page")

            if not text:
                continue

            start = 0
            text_length = len(text)

            while start < text_length:

                end = min(
                    start + self.chunk_size,
                    text_length,
                )

                chunk_text = text[start:end].strip()

                if chunk_text:

                    chunks.append(
                        {
                            "id": (
                                f"{source_name}:"
                                f"{chunk_id}"
                            ),
                            "text": chunk_text,
                            "source": source_name,
                            "page": page,
                            "chunk_index": chunk_id,
                        }
                    )

                    chunk_id += 1

                if end >= text_length:
                    break

                start = end - self.chunk_overlap

        return chunks